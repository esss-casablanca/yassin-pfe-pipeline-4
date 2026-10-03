#!/usr/bin/env python3
"""build_deck.py — build the ESSS soutenance deck from a slide specification (Pipeline 4, skills p4-01 / p4-02).

Usage:
  python build_deck.py --spec deck_spec.json --template esss_soutenance_template.pptx \
      --out soutenance_deck_V1.pptx [--notes-docx soutenance_notes_V1.docx] [--timing timing_V1.json]

deck_spec.json
{
  "meta": {"academic_year": "2025/2026", "footer": "Soutenance PFE · ESSS · P23",
           "presentation_minutes": 20},
  "slides": [
    {"type": "title",    "title": "…", "subtitle": "Prénom NOM · Filière · Promotion · Encadrant : …", "notes": "…", "seconds": 30},
    {"type": "section",  "title": "Problématique et question de recherche", "subtitle": "…", "seconds": 10},
    {"type": "bullets",  "title": "…", "bullets": ["…", {"text": "…", "level": 1}], "notes": "…", "seconds": 50},
    {"type": "two_columns", "title": "…", "left_label": "Composante 1", "left": ["…"], "right_label": "Composante 2", "right": ["…"]},
    {"type": "table",    "title": "…", "header": ["…"], "rows": [["…"]], "font_pt": 14, "caption": "Tableau 2 …", "col_widths": [3, 1, 1]},
    {"type": "image",    "title": "…", "image": "prisma_flow.png", "caption": "Figure 1 …"},
    {"type": "chart",    "title": "…", "chart": {"kind": "column|bar|pie|line", "categories": ["…"],
                          "series": [{"name": "…", "values": [1, 2]}], "number_format": "0.0", "legend": true}},
    {"type": "closing",  "title": "Merci de votre attention", "subtitle": "Prénom NOM — questions"},
    {"type": "annex_marker", "title": "Annexes"},        # everything after it is an annex (no timing)
    {"type": "table", "annex": true, …}
  ]
}
Every slide accepts "notes" (speaker notes, French) and "seconds" (target speaking time; annex slides count 0).
Lint warnings (too many bullets, long bullets, generic titles, missing notes) go to the timing JSON and stdout.
"""
from __future__ import annotations

import argparse
import copy
import json
import os
import re
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Emu, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

LAYOUT = {"title": 0, "content": 1, "section": 2, "two": 3, "comparison": 4, "title_only": 5, "blank": 6}
DEFAULT_SECONDS = {"title": 30, "section": 10, "bullets": 50, "two_columns": 55, "table": 60, "image": 50,
                   "chart": 50, "closing": 10, "annex_marker": 0}
GENERIC_TITLES = re.compile(r"^\s*(résultats?|resultats?|discussion|méthodes?|methodes?|introduction|conclusion)\s*\d*\s*$", re.I)

CONTENT_LEFT, CONTENT_TOP, CONTENT_W, CONTENT_H = 581192, 2180496, 11029615, 3678303


# ------------------------------------------------------------------------------------------------
def clear_slides(prs: Presentation) -> None:
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        prs.part.drop_rel(sldId.rId)
        sldIdLst.remove(sldId)


def set_year_on_title_layout(prs: Presentation, academic_year: str) -> None:
    layout = prs.slide_layouts[LAYOUT["title"]]
    for sh in layout.shapes:
        if sh.has_text_frame and "ANNEE UNIVERSITAIRE" in sh.text_frame.text.upper():
            pars = sh.text_frame.paragraphs
            wanted = ["PROJET DE FIN D’ÉTUDES", f"ANNÉE UNIVERSITAIRE {academic_year}"]
            for p, txt in zip(pars, wanted):
                runs = p.runs
                if runs:
                    runs[0].text = txt
                    for r in runs[1:]:
                        r._r.getparent().remove(r._r)
                else:
                    p.text = txt


def add_footer_and_number(slide, layout, footer_text: str) -> None:
    """Clone the layout's footer (idx 11) and slide-number (idx 12) placeholders onto the slide."""
    for ph in layout.placeholders:
        idx = ph.placeholder_format.idx
        if idx not in (11, 12):
            continue
        el = copy.deepcopy(ph._element)
        slide.shapes._spTree.append(el)
    for sh in slide.placeholders:
        if sh.placeholder_format.idx == 11:
            sh.text_frame.text = footer_text


def fill_bullets(text_frame, bullets, max_level=4, font_pt: int | None = None) -> None:
    text_frame.clear()
    first = True
    for b in bullets:
        if isinstance(b, str):
            text, level = b, 0
        else:
            text, level = b.get("text", ""), int(b.get("level", 0))
        p = text_frame.paragraphs[0] if first else text_frame.add_paragraph()
        first = False
        p.text = text
        p.level = max(0, min(level, max_level))
        if font_pt:
            for r in p.runs:
                r.font.size = Pt(font_pt - 2 * p.level)


def add_notes(slide, notes: str) -> None:
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def add_table(slide, spec, top=CONTENT_TOP, height=CONTENT_H) -> None:
    header = spec.get("header") or []
    rows = spec.get("rows") or []
    n_rows = len(rows) + (1 if header else 0)
    n_cols = max([len(header)] + [len(r) for r in rows]) if (header or rows) else 1
    shape = slide.shapes.add_table(n_rows, n_cols, Emu(CONTENT_LEFT), Emu(top), Emu(CONTENT_W), Emu(min(height, 400000 * n_rows + 200000)))
    table = shape.table
    widths = spec.get("col_widths")
    if widths and len(widths) == n_cols:
        total = float(sum(widths))
        for j, w in enumerate(widths):
            table.columns[j].width = Emu(int(CONTENT_W * w / total))
    font_pt = int(spec.get("font_pt", 14))
    r0 = 0
    if header:
        for j, h in enumerate(header):
            cell = table.cell(0, j)
            cell.text = str(h)
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(font_pt)
                    r.font.bold = True
        r0 = 1
    for i, row in enumerate(rows):
        for j in range(n_cols):
            val = row[j] if j < len(row) else ""
            cell = table.cell(r0 + i, j)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(font_pt)
    if spec.get("caption"):
        tb = slide.shapes.add_textbox(Emu(CONTENT_LEFT), Emu(CONTENT_TOP + CONTENT_H - 320000), Emu(CONTENT_W), Emu(320000))
        tb.text_frame.text = spec["caption"]
        for p in tb.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(12)
                r.font.italic = True


def add_image(slide, spec) -> None:
    path = spec["image"]
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    from PIL import Image  # pillow ships with python-pptx installs on most systems
    with Image.open(path) as im:
        w, h = im.size
    box_w, box_h = CONTENT_W, CONTENT_H - (350000 if spec.get("caption") else 0)
    scale = min(box_w / w, box_h / h)
    pw, ph = int(w * scale), int(h * scale)
    left = CONTENT_LEFT + (box_w - pw) // 2
    slide.shapes.add_picture(path, Emu(left), Emu(CONTENT_TOP), Emu(pw), Emu(ph))
    if spec.get("caption"):
        tb = slide.shapes.add_textbox(Emu(CONTENT_LEFT), Emu(CONTENT_TOP + box_h + 20000), Emu(CONTENT_W), Emu(320000))
        tb.text_frame.text = spec["caption"]
        for p in tb.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(12)
                r.font.italic = True


def add_chart(slide, spec) -> None:
    c = spec["chart"]
    kind = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
            "pie": XL_CHART_TYPE.PIE, "line": XL_CHART_TYPE.LINE_MARKERS}[c.get("kind", "column")]
    data = CategoryChartData()
    data.categories = c["categories"]
    for s in c["series"]:
        data.add_series(s["name"], s["values"])
    gf = slide.shapes.add_chart(kind, Emu(CONTENT_LEFT), Emu(CONTENT_TOP), Emu(CONTENT_W), Emu(CONTENT_H), data)
    chart = gf.chart
    chart.has_legend = bool(c.get("legend", len(c["series"]) > 1))
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.number_format = c.get("number_format", "0.0")
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.font.size = Pt(12)
    if c.get("kind", "column") in ("column", "bar", "line"):
        try:
            chart.value_axis.minimum_scale = float(c.get("axis_min", 0))
            if c.get("axis_max") is not None:
                chart.value_axis.maximum_scale = float(c["axis_max"])
            chart.value_axis.has_major_gridlines = True
        except Exception:
            pass
    if c.get("title"):
        chart.has_title = True
        chart.chart_title.text_frame.text = c["title"]


# ------------------------------------------------------------------------------------------------
def build(spec: dict, template: str, out: str, notes_docx: str | None, timing_path: str | None) -> dict:
    prs = Presentation(template)
    clear_slides(prs)
    meta = spec.get("meta", {})
    set_year_on_title_layout(prs, meta.get("academic_year", ""))
    footer = meta.get("footer", "Soutenance PFE · ESSS")

    timing, warnings = [], []
    in_annex = False
    for n, s in enumerate(spec["slides"], 1):
        t = s["type"]
        if t == "annex_marker":
            in_annex = True
        annex = in_annex or bool(s.get("annex"))
        if t == "title":
            layout = prs.slide_layouts[LAYOUT["title"]]
            slide = prs.slides.add_slide(layout)
            slide.shapes.title.text = s.get("title", "")
            slide.placeholders[1].text = s.get("subtitle", "")
        elif t in ("section", "annex_marker"):
            layout = prs.slide_layouts[LAYOUT["section"]]
            slide = prs.slides.add_slide(layout)
            slide.shapes.title.text = s.get("title", "Annexes" if t == "annex_marker" else "")
            slide.placeholders[1].text = s.get("subtitle", "")
        elif t == "bullets":
            layout = prs.slide_layouts[LAYOUT["content"]]
            slide = prs.slides.add_slide(layout)
            slide.shapes.title.text = s.get("title", "")
            fill_bullets(slide.placeholders[1].text_frame, s.get("bullets", []), font_pt=s.get("font_pt"))
        elif t == "two_columns":
            labelled = bool(s.get("left_label") or s.get("right_label"))
            layout = prs.slide_layouts[LAYOUT["comparison"] if labelled else LAYOUT["two"]]
            slide = prs.slides.add_slide(layout)
            slide.shapes.title.text = s.get("title", "")
            fp = s.get("font_pt", 18)
            if labelled:
                slide.placeholders[1].text = s.get("left_label", "")
                fill_bullets(slide.placeholders[2].text_frame, s.get("left", []), font_pt=fp)
                slide.placeholders[3].text = s.get("right_label", "")
                fill_bullets(slide.placeholders[4].text_frame, s.get("right", []), font_pt=fp)
            else:
                fill_bullets(slide.placeholders[1].text_frame, s.get("left", []), font_pt=fp)
                fill_bullets(slide.placeholders[2].text_frame, s.get("right", []), font_pt=fp)
        elif t == "table":
            layout = prs.slide_layouts[LAYOUT["title_only"]]
            slide = prs.slides.add_slide(layout)
            slide.shapes.title.text = s.get("title", "")
            add_table(slide, s)
        elif t == "image":
            layout = prs.slide_layouts[LAYOUT["title_only"]]
            slide = prs.slides.add_slide(layout)
            slide.shapes.title.text = s.get("title", "")
            add_image(slide, s)
        elif t == "chart":
            layout = prs.slide_layouts[LAYOUT["title_only"]]
            slide = prs.slides.add_slide(layout)
            slide.shapes.title.text = s.get("title", "")
            add_chart(slide, s)
        elif t == "closing":
            layout = prs.slide_layouts[LAYOUT["section"]]
            slide = prs.slides.add_slide(layout)
            slide.shapes.title.text = s.get("title", "Merci de votre attention")
            slide.placeholders[1].text = s.get("subtitle", "")
        else:
            raise SystemExit(f"slide {n}: unknown type {t!r}")

        if t != "title":
            add_footer_and_number(slide, layout, footer)
        add_notes(slide, s.get("notes", ""))

        secs = 0 if annex else int(s.get("seconds", DEFAULT_SECONDS.get(t, 45)))
        timing.append({"slide": n, "type": t, "title": s.get("title", ""), "seconds": secs, "annex": annex})

        # lint
        if t in ("bullets", "two_columns"):
            bl = s.get("bullets", []) + s.get("left", []) + s.get("right", [])
            if len(s.get("bullets", [])) > 6:
                warnings.append(f"slide {n}: {len(s['bullets'])} bullets (> 6)")
            for b in bl:
                txt = b if isinstance(b, str) else b.get("text", "")
                if len(txt.split()) > 14:
                    warnings.append(f"slide {n}: long bullet ({len(txt.split())} words): {txt[:60]}…")
        if t not in ("title", "section", "closing", "annex_marker") and GENERIC_TITLES.match(s.get("title", "")):
            warnings.append(f"slide {n}: generic title {s.get('title')!r} — make it a question or a statement")
        if t not in ("section", "annex_marker", "closing") and not annex and not s.get("notes"):
            warnings.append(f"slide {n}: no speaker notes")

    prs.save(out)

    total = sum(x["seconds"] for x in timing)
    target = int(meta.get("presentation_minutes", 20)) * 60
    result = {"deck": os.path.basename(out), "slides": len(timing), "timed_slides": sum(1 for x in timing if not x["annex"]),
              "annex_slides": sum(1 for x in timing if x["annex"]), "estimated_seconds": total,
              "estimated_minutes": round(total / 60, 1), "target_minutes": target // 60,
              "within_target": abs(total - target) <= 0.1 * target, "per_slide": timing, "warnings": warnings}
    if timing_path:
        with open(timing_path, "w", encoding="utf-8") as f:
            json.dump(result, f, ensure_ascii=False, indent=2)
    if notes_docx:
        write_notes_docx(spec, timing, result, notes_docx)
    return result


def write_notes_docx(spec: dict, timing: list[dict], result: dict, path: str) -> None:
    import docx
    from docx.shared import Pt as DPt
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = DPt(11)
    meta = spec.get("meta", {})
    d.add_heading("Script de soutenance — notes de l’orateur", 0)
    d.add_paragraph(f"{meta.get('footer', '')} · durée cible {result['target_minutes']} min · "
                    f"estimation {result['estimated_minutes']} min sur {result['timed_slides']} diapositives (+ {result['annex_slides']} annexes)")
    running = 0
    for s, tm in zip(spec["slides"], timing):
        running += tm["seconds"]
        head = f"Diapositive {tm['slide']} — {s.get('title', '') or s['type']}"
        d.add_heading(head, level=2)
        if tm["annex"]:
            d.add_paragraph("[annexe — hors chronométrage]").italic = True
        else:
            d.add_paragraph(f"[⏱ {tm['seconds']} s · cumul {running // 60} min {running % 60:02d} s]").italic = True
        for para in (s.get("notes", "") or "(pas de notes)").split("\n"):
            if para.strip():
                d.add_paragraph(para.strip())
    d.add_heading("Chronométrage", level=1)
    tbl = d.add_table(rows=1, cols=4)
    tbl.style = "Light Grid Accent 1"
    hdr = tbl.rows[0].cells
    for c, t in zip(hdr, ("N°", "Diapositive", "Secondes", "Cumul")):
        c.text = t
    running = 0
    for tm in timing:
        running += tm["seconds"]
        row = tbl.add_row().cells
        row[0].text = str(tm["slide"])
        row[1].text = (tm["title"] or tm["type"])[:70]
        row[2].text = "annexe" if tm["annex"] else str(tm["seconds"])
        row[3].text = f"{running // 60}:{running % 60:02d}"
    d.save(path)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--spec", required=True)
    ap.add_argument("--template", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--notes-docx")
    ap.add_argument("--timing")
    a = ap.parse_args()
    with open(a.spec, encoding="utf-8") as f:
        spec = json.load(f)
    res = build(spec, a.template, a.out, a.notes_docx, a.timing)
    print(f"deck: {res['deck']} — {res['slides']} slides ({res['timed_slides']} timed + {res['annex_slides']} annex); "
          f"estimated {res['estimated_minutes']} min for a {res['target_minutes']} min target "
          f"({'OK' if res['within_target'] else 'ADJUST'})")
    for w in res["warnings"]:
        print("  warn:", w)
    return 0


if __name__ == "__main__":
    sys.exit(main())
