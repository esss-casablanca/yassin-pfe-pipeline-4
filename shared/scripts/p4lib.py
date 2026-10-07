#!/usr/bin/env python3
"""p4lib — shared helpers for the Pipeline-4 scripts (Python 3.10+, python-docx 1.2, python-pptx 1.0).

Number normalisation is locale-aware:
  fr : thousands separated by (narrow/no-break) spaces, decimal comma   -> "1 248" "12,5"
  en : thousands separated by commas, decimal point                      -> "1,248" "12.5"
  ar : ASCII digits by Pipeline-4 contract, decimal point; Arabic-Indic digits are mapped to ASCII first
A normalised number is a plain string such as "1248", "12.5", "-0.03".

Parser version 2 (plugin v0.2.4) — fixes reported on 2026-10-04 (D15-P04):
  * thousands separators: en/ar also accept no-break / narrow / thin spaces and the Arabic separator U+066C;
    the Arabic decimal separator U+066B is read as a decimal point;
  * bracketed intervals ("IC 95 % [1,3 ; 3,4]", "28 [24–33]") are read as numbers, no longer as citation keys;
  * DOIs and URLs in the running text are ignored;
  * section cross-references ("2.3", "§ 2.3", "القسم 2.3", "2.3.1") are read identically in every locale
    (French accepts a decimal point as a fallback, multi-level numbers are ignored, Arabic section words are labels);
  * a French label followed by a grouped number ("tableau 3 120 femmes") is split into the label and the value.
A ledger built with an older parser must be rebuilt (extract_ledger.py --carry-over) before reconciliation.

Parser version 4 (plugin v0.2.7) — fixes reported on 2026-10-06 (D10-P01):
  * Unicode bidi controls (LRE/PDF/LRM/RLM/isolates, zero-width space, BOM) are stripped before any matching, so an
    Arabic paragraph whose numeric range was embedded left-to-right by the typesetter still reads "[29–33]" as a
    citation and not as two introduced numbers;
  * the reference list ends at the first paragraph that is not a reference entry (an "Annexe A" title typed without
    a heading style, an acknowledgement, a table …), not only at the next styled heading (walk_blocks);
  * numbered section headings ("7. Références", "8. Annexes") and the English/Arabic variants of the headings are
    recognised; body-level content controls (Word's bibliography) and text boxes are read;
  * ordinals ("1er octobre", "2e", "3rd") are read as the number they carry; a list number at the start of a paragraph
    ("1. …", "2) …") is tagged role_hint "numbering" and ignored by the reconciliation on both sides;
  * in Arabic, a digit glued to a clitic (و78,8 %، الـ210، ب2) is read as the number.
"""
from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from dataclasses import dataclass, field, asdict
from typing import Iterable

PARSER_VERSION = "5"  # bump whenever find_numbers / find_citations change what they return
PARSER_COMPATIBLE = {"3", "4"}  # older ledgers the reconciler still accepts (with a warning) — see reconcile_ledger.py

GROUP_SPACES = "     "  # nbsp, narrow nbsp, thin space, figure space, space
EN_GROUP = ",    ٬"  # comma, no-break spaces, Arabic thousands separator (not the plain space)
AR_DECIMAL = "٫"  # Arabic decimal separator
ARABIC_INDIC = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789")
DASHES = "–—−-"  # en dash, em dash, minus, hyphen
# Unicode format characters that carry no text: bidi embeddings/overrides/isolates and marks (U+202A–E, U+2066–9,
# U+200E/F), the zero-width space and the byte-order mark. A typesetter may insert them (an LTR embedding keeps
# "29–33" readable inside an Arabic line); they must never change what the parser reads.
BIDI_CONTROLS = re.compile("[\u200b\u200e\u200f\u202a-\u202e\u2066-\u2069\ufeff\u00ad\ue000-\uf8ff]")
# (also the soft hyphen and the private-use glyphs Word writes for Symbol/Wingdings bullets, e.g. U+F0B7 — v0.2.7)
# a bullet typed by hand in front of a reference entry ("• Dusse F,", "- Soufi G,", "o Dawe, J.")
_BULLET = r"(?:[•▪■◦‣·\-–—*]+\s*|o\s+)?"
LEAD_BULLET = re.compile(r"^\s*(?:[•▪■◦‣·\-–—*]+\s*|o\s+)")

CITATION_NUMERIC = re.compile(r"\[(\d{1,3}(?:\s*[,;–—-]\s*\d{1,3})*)\]")
_NAME_WORD = r"[A-ZÀ-ÝŒ](?:[^\W\d_]|['’\-](?=[^\W\d_]))+"      # "Dall’Ora", "Crea-Arsenio", "O’Brien" — never "COVID-19"
_NAME = rf"{_NAME_WORD}(?:\s+(?:&|and|et|و)\s+{_NAME_WORD})*(?:\s+(?:et\s+al\.?|و\s?آخرون))?"
CITATION_AUTHOR_DATE = re.compile(
    rf"\(((?:{_NAME})\s*,?\s*(?:19|20)\d{{2}}[a-z]?(?:\s*;\s*(?:{_NAME})\s*,?\s*(?:19|20)\d{{2}}[a-z]?)*)\)"
)
# narrative author–date citation: "Williams et al. (2023)", "Aiken & Clarke (2002)", "Williams وآخرون (2023)" (v0.2.7)
CITATION_NARRATIVE = re.compile(rf"(?<![\w'’\-])({_NAME})\s*\(\s*((?:19|20)\d{{2}}[a-z]?)\s*\)")
_CIT_ETAL = re.compile(r"\s*(?:\bet\s+al\.?|و\s?آخرون)\s*", re.I)
_CIT_AND = re.compile(r"\s+(?:&|and|et|و)\s+(?!al\b)")

_SECTION_NO = r"(?:\d{1,2}[.)]?\s*)?"  # an optional section number before a heading ("7. Références")
HEADING_REFS = re.compile(
    rf"^\s*{_SECTION_NO}(r[ée]f[ée]rences?(\s+bibliographiques)?|bibliographie|bibliography|references?(\s+list)?|"
    r"liste\s+des\s+r[ée]f[ée]rences|works\s+cited|(?:قائمة\s+)?المراجع(?:\s+والمصادر)?|المصادر(?:\s+والمراجع)?)\s*:?\s*$", re.I)
# a heading that opens with the word ("Références (complètes — 43 études)", "References (verified)", "B. Références de
# contexte") or a short heading that contains it ("Références retirées") also opens the reference list (v0.2.7)
HEADING_REFS_LOOSE = re.compile(r"^\s*(?:[A-Za-z]|\d{1,2})?[.)]?\s*(r[ée]f[ée]rences?|references?|bibliograph\w*|(?:قائمة\s+)?المراجع)\b", re.I)
# "Sources et références", "Bibliographie / Références", "Liste de références" — the word closes a short heading
# without figures (a checklist line "… (excluding abstract/references)" or "2.4 Sources d'information" is not one)
_REFS_WORD = re.compile(r"^\s*(?:[A-Za-z]|\d{1,2})?[.)]?\s*[^\d()]{0,30}?(?<![\w])(r[ée]f[ée]rences?|references?|bibliograph\w*|المراجع)\s*:?\s*$", re.I)
HEADING_ANNEX = re.compile(rf"^\s*{_SECTION_NO}(annexes?|appendix|appendices|supplementary|suppl[ée]ments?|الملاحق|الملحق|ملاحق|ملحق)\b", re.I)


def is_refs_heading(text: str) -> bool:
    """True for a heading that opens the reference list."""
    t = (text or "").strip()
    if not t or HEADING_ANNEX.match(t):
        return False
    return bool(HEADING_REFS.match(t) or HEADING_REFS_LOOSE.match(t) or (len(t) <= 60 and _REFS_WORD.match(t)))
# Paragraphs that open a section placed after the reference list (typed without a heading style, they would be
# swallowed as reference entries): annexes, acknowledgements, declarations, funding, contributions, captions …
AFTER_REFS_SECTION = re.compile(
    rf"^\s*{_SECTION_NO}(annexes?|appendix|appendices|supplementary|suppl[ée]ments?|mat[ée]riels?\s+suppl\w*|"
    r"donn[ée]es\s+suppl\w*|remerciements?|acknowledg\w*|"
    r"d[ée]clarations?|declarations?|conflits?\s+d|conflicts?\s+of|financement|funding|contributions?|"
    r"disponibilit[ée]\s+des\s+donn[ée]es|data\s+availability|tableau|table|figure|fig\.|encadr[ée]|"
    r"الملاحق|الملحق|ملاحق|ملحق|شكر|إقرار|تصريح|تمويل|مساهم\w*|الجدول|جدول|الشكل|شكل)\b", re.I)
REF_INDEX = re.compile(r"^\s*" + _BULLET + r"\[?(\d{1,3})([a-d])?\]?(?:[.)]\s*|\s+)(?=[^\d\s])")
# "12. Author", "[12] Author", "18.Price" (no space), "[5b] Author" (sub-indexed entry), "• 3. Author" (hand bullet)
NOTE_OPENER = re.compile(r"^\s*(?:note|nb|n\.b\.|remarque|nota|ملاحظة|تنبيه)\b", re.I)
# how a reference entry opens: "Surname AB," / "Surname AB." / "van der Heijden B," / "Organisation name." / "World
# Health Organization (WHO)." — a capitalised word (after an optional particle) followed by initials, or a run of
# capitalised words ending with a period
_PARTICLES = r"(?:(?i:van|von|de|der|den|du|da|di|le|la|el|al|bin|ibn|ben|mac|mc|o'|d'|l')[\s'’]*)*"
_SURNAME = r"[A-ZÀ-ÝŒ][\w'’\-À-ÿ]*(?:[\s-]+(?:" + _PARTICLES.rstrip("*") + r")?[A-ZÀ-ÝŒ][\w'’\-À-ÿ]*)*"
REF_AUTHOR_START = re.compile(
    r"^\s*" + _BULLET + r"(?:\[?\d{1,3}[a-d]?\]?(?:[.)]\s*|\s+))?" + _PARTICLES + r"(?:"
    + _SURNAME + r"\s+[A-ZÀ-ÝŒ]{1,4}\b(?:[.,;]|\s+(?:et\s+al|and|&|وآخرون))"   # Vancouver: "Soufi G,", "Vo V et al. (2023)"
    + r"|" + _SURNAME + r",\s*(?:[A-ZÀ-ÝŒ]\.?(?:\s*-?\s*[A-ZÀ-ÝŒ]\.?)*)[,.;&\s]"  # APA: "Dawe, J.,", "de Moissac, D., &"
    + r"|[A-ZÀ-ÝŒ][\w'’\-À-ÿ]+(?:\s+(?:[A-ZÀ-ÝŒ][\w'’\-À-ÿ]+|de|du|des|la|le|et|of|and|the|for|على|في|و))*\s*(?:\([A-Z]+\))?"
    + r"\s*(?:\.|\(\s*(?:19|20)\d{2}[a-z]?\s*\)))"                           # organisation: "WHO.", "WHO (2021)."
)
TABLE_LABEL = re.compile(r"^\s*(tableau|table|الجدول|جدول)\s*(\d+[a-z]?)\b", re.I)
FIGURE_LABEL = re.compile(r"^\s*(figure|fig\.|الشكل|شكل)\s*(\d+[a-z]?)\b", re.I)


def sha256_file(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def nfc(s: str) -> str:
    """NFC normalisation; bidi controls, zero-width spaces, BOMs, soft hyphens and private-use bullet glyphs are
    removed (they are formatting, not text)."""
    return BIDI_CONTROLS.sub("", unicodedata.normalize("NFC", s or ""))


def strip_bullet(text: str) -> str:
    """The entry without a hand-typed bullet in front of it."""
    return LEAD_BULLET.sub("", (text or "").strip())


def split_ref_index(text: str) -> tuple[int | None, str]:
    """(index, body) of a reference entry: the index typed in front of it ("12.", "[12]", "[5b]" → 5) or None, and
    the entry without index or bullet. One reader for the ledger, the carry-over and the reconciliation (v0.2.7)."""
    t = strip_bullet(text)
    m = REF_INDEX.match(t)
    if not m:
        return None, t
    return int(m.group(1)), t[m.end():].strip()


# ---------------------------------------------------------------------------------------------
# Number extraction
# ---------------------------------------------------------------------------------------------
@dataclass
class NumberHit:
    raw: str
    normalized: str
    start: int
    end: int
    unit: str | None = None
    role_hint: str | None = None
    context: str = ""


# ordinal suffixes glued to a number ("1er octobre", "2e", "3rd"): the number is read, the suffix is not a word
_ORDINAL = r"(?P<ord>(?:er|ère|ere|re|e|ème|eme|st|nd|rd|th))?"
# Letters typed against a number never hide it (v0.2.8). Students lose the spaces of table cells ("IC95%", "n45",
# "p0,03", "45ans", "150patients", "OR2,1", "mixtes115", "etal.,2021Mondial") and the translation puts them back
# ("95% CI", "n = 45", "p = 0.03", "45 years"): the two sides must read the same figures, so the parser reads the
# digits whatever letter precedes or follows them — codes and identifiers included ("D03-P04" reads 3 and 4,
# "H1N1" 1 and 1, "B12" 12): they are copied as typed on both sides, so they compare equal, and the ledger simply
# carries them. What still hides a number: a digit, "_" or a decimal/section separator right before it ("2,5"
# never reads 5, "2.3.1" is blanked), an item/citation label glued after it ("14b", "2025b", "23a-d",
# "Tableau 8bis"), an address ("dbali45@gmail.com"); a hexadecimal hash (SHA-256) and a URL/DOI are
# blanked beforehand; the "-" of "pre-2015", "Round-2" or "WHO-5" is a hyphen, never a minus sign. A decimal typed
# without its zero after a space, "=", "<", "(" or "[" (APA "p = .05", ",009") reads 0.05 and 0.009.
_UP = "A-ZÀ-ÖØ-ÝŒ"          # Latin capitals (× U+00D7 excluded: "2×2" reads two numbers)
_BEFORE_NUM = r"(?<![0-9_])(?<!\d[,.])"
_SIGN = r"(?:(?<![\w])[-−])?"
# a capital glued after the number does not hide it either ("2023PRFI", "2D03-P04", "2021Mondial" read 2023, 2, 2021:
# a lost space before an acronym is far more common in a deposited table than a code like "6S", which reads 6 on
# both sides anyway); an item/citation label does ("14b", "8bis")
_AFTER_NUM = r"(?![\d_@])(?!(?:[a-d]|bis|ter)(?![a-zà-ÿ]))"
# ".05" / ",009": a leading decimal separator after a space, "=", "<", "(" or "[" (or at the start of the text)
_LEAD_DEC = r"(?:(?<=[=<>≤≥(:;/\[])|(?<![^\s]))[.,]\d+"
HEX_HASH = re.compile(r"\b[0-9a-f]{32,}\b")   # SHA-256 and the like, copied verbatim on both sides


def _fr_pattern() -> re.Pattern:
    sp = f"[{GROUP_SPACES}]"
    # decimal comma (French) or, as a fallback, decimal point: "2.3" must read the same in every locale
    return re.compile(
        rf"{_BEFORE_NUM}(?P<num>{_SIGN}(?:\d{{1,3}}(?:{sp}\d{{3}})+|\d+)(?:[,.]\d+)?|{_LEAD_DEC}){_ORDINAL}{_AFTER_NUM}"
    )


def _en_pattern() -> re.Pattern:
    g = f"[{EN_GROUP}]"
    return re.compile(rf"{_BEFORE_NUM}(?P<num>{_SIGN}(?:\d{{1,3}}(?:{g}\d{{3}})+|\d+)(?:[.{AR_DECIMAL}]\d+)?|{_LEAD_DEC}){_ORDINAL}{_AFTER_NUM}")


def _ar_pattern() -> re.Pattern:
    # Arabic orthography glues clitics to the next word: "و78,8 %" (and 78.8 %), "الـ210" (the 210), "ب2" (with 2).
    # An Arabic letter or the tatweel before a digit therefore does not hide the number; a Latin letter or a digit does.
    g = f"[{EN_GROUP}]"
    return re.compile(rf"{_BEFORE_NUM}(?P<num>{_SIGN}(?:\d{{1,3}}(?:{g}\d{{3}})+|\d+)(?:[.{AR_DECIMAL}]\d+)?|{_LEAD_DEC}){_AFTER_NUM}")


_PATTERNS = {"fr": _fr_pattern(), "en": _en_pattern(), "ar": _ar_pattern()}


def normalize_number(raw: str, locale: str) -> str:
    s = raw.translate(ARABIC_INDIC).replace("−", "-")
    if locale == "fr":
        for ch in GROUP_SPACES:
            s = s.replace(ch, "")
        s = s.replace(",", ".")
    else:  # en / ar
        for ch in EN_GROUP:
            s = s.replace(ch, "")
        s = s.replace(AR_DECIMAL, ".")
    # canonical string: no leading zeros on the integer part; written precision kept (12.50 stays 12.50)
    neg = s.startswith("-")
    s = s.lstrip("-")
    if "." in s:
        ip, fp = s.split(".", 1)
        s = (ip.lstrip("0") or "0") + "." + fp
    else:
        s = s.lstrip("0") or "0"
    return ("-" if neg else "") + s


# ---- things that look like numbers but are not values of the paper -----------------------------------------
URL_DOI = re.compile(r"(?i)(?:https?://|www\.)\S+|\bdoi\s*:\s*\S+|\b10\.\d{4,9}/[^\s\"<>]+"
                     r"|\b[\w-]+(?:\.[\w-]+)*\.(?:com|org|net|app|edu|gov|int|io|ma|fr|ca|uk|ch|be|eu|info)/\S+")  # bare "consensus.app/…"
# numeric dates (06/10/2026, 6-10-2026, 2026-10-06): read as one item of role "date", verified by hand in the
# translations, where the format legitimately changes ("6 October 2026") — v0.2.7
DATE_NUMERIC = re.compile(r"(?<![\w.,])(?:(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})|(\d{1,2})[-/.](\d{1,2})[-/.](\d{4}))(?![\w])")
MONTHS = {
    1: "janvier|janv\\.?|january|jan\\.?|يناير|كانون الثاني", 2: "f[ée]vrier|f[ée]v\\.?|february|feb\\.?|فبراير|شباط",
    3: "mars|march|mar\\.?|مارس|آذار", 4: "avril|avr\\.?|april|apr\\.?|أبريل|إبريل|ابريل|نيسان",
    5: "mai|may|ماي|مايو|أيار", 6: "juin|june|jun\\.?|يونيو|يونيه|حزيران", 7: "juillet|juil\\.?|july|jul\\.?|يوليو|يوليوز|تموز",
    8: "ao[uû]t|august|aug\\.?|أغسطس|غشت|آب", 9: "septembre|sept?\\.?|september|سبتمبر|شتنبر|أيلول",
    10: "octobre|oct\\.?|october|أكتوبر|اكتوبر|تشرين الأول", 11: "novembre|nov\\.?|november|نوفمبر|نونبر|تشرين الثاني",
    12: "d[ée]cembre|d[ée]c\\.?|december|dec\\.?|ديسمبر|دجنبر|كانون الأول",
}
_MONTH_ALT = "|".join(f"(?P<m{k}>{v})" for k, v in MONTHS.items())
# "6 octobre 2025", "1er octobre 2025", "October 6, 2025", "6 أكتوبر 2025": one date item, like the numeric form
DATE_WORDS = re.compile(
    rf"(?<![\w.,])(?:(?P<d1>\d{{1,2}})(?:er|ère|re|e|st|nd|rd|th)?\s+(?:{_MONTH_ALT})\s+(?P<y1>(?:19|20)\d{{2}})"
    rf"|(?:{_MONTH_ALT.replace('(?P<m', '(?P<n')})\s+(?P<d2>\d{{1,2}})(?:st|nd|rd|th)?,?\s+(?P<y2>(?:19|20)\d{{2}}))(?![\w])", re.I)
MULTI_DOT = re.compile(r"(?<![\w.,])\d+(?:\.\d+){2,}(?![\w])")  # 2.3.1 — section / item numbering
LABEL_BEFORE = re.compile(
    r"(?i)(tableau|table|figure|fig\.|الجدول|الشكل|جدول|شكل|composante|component|étape|step|phase|partie|part|"
    r"section|chapitre|chapter|annexe|appendix|objectif|objective|hypoth[èe]se|hypothesis|paragraphe|paragraph|§|"
    r"المكو[نّ]+|الملحق|المرحلة|القسم|الفصل|الجزء|الفقرة|البند|الباب)\s*$")

# ---- bracketed intervals versus numeric citation keys ---------------------------------------------------------
_STAT_ACRONYMS = re.compile(r"(?:\b(?:IC|CI|OR|aOR|ORa|ORaj|RR|aRR|RRa|HR|aHR|RP|PR|IRR|EIQ|IIQ|IQR|SD|ET|DS)\b|95\s*%|±)")
# v0.2.7: Arabic words with or without the article (وسيط العمر / الوسيط), the interquartile range, mean ± SD
_STAT_WORDS = re.compile(r"(?i)(?:intervalle|interval|interquartile|étendue|\brange\b|m[ée]diane?\b|\bmin(?:imum)?\b|"
                         r"\bmax(?:imum)?\b|moyenne|\bmean\b|average|écart[- ]type|standard\s+deviation|"
                         r"فترة\s+(?:ال)?ثقة|فاصل\s+(?:ال)?ثقة|مجال\s+(?:ال)?ثقة|(?:ال)?مدى|(?:ال)?وسيط|(?:ال)?متوسط|"
                         r"(?:ال)?ربيعي|نسبة\s+(?:ال)?أرجحية|الحد\s+الأدنى|الحد\s+الأقصى|أدنى|أقصى|انحراف)")
_NUMBER_JUST_BEFORE = re.compile(r"(?<![\w.,])(\d+(?:[,.]\d+)?)\s*%?\s*$")


def _stat_context(before: str) -> bool:
    """Statistical vocabulary in the 40 characters before a bracket, within the same clause."""
    w = before[-40:]
    for stop in ("]", ". ", "\n"):
        k = w.rfind(stop)
        if k >= 0:
            w = w[k + len(stop):]
    return bool(_STAT_ACRONYMS.search(w) or _STAT_WORDS.search(w))


def plausible_citation_list(inner: str) -> bool:
    """True when the bracket content reads as a well-formed numeric citation list: integer keys and ranges a–b
    (a < b, span ≤ 20), strictly ascending, never repeated — "1,3-5,7" and "12–15" yes; "52,5–79,5" (52 then 5),
    "2,6–51,3" (span 45) and "30,1–49,5" (30 then 1) no: those are decimal-comma intervals."""
    last = 0
    for part in re.split(r"\s*[,;]\s*", inner.strip()):
        if not part:
            return False
        m = re.fullmatch(r"(\d{1,3})\s*[–—-]\s*(\d{1,3})", part)
        if m:
            a, b = int(m.group(1)), int(m.group(2))
            if a >= b or b - a > 20:
                return False
        elif re.fullmatch(r"\d{1,3}", part):
            a = b = int(part)
        else:
            return False
        if a <= last:
            return False
        last = b
    return True


def bracket_is_interval(inner: str, before: str) -> bool:
    """True when a bracket matched by CITATION_NUMERIC is a numeric interval (CI, IQR, range), not citation keys.

    Citation keys are positive integers without leading zeros. An interval has two bounds: two decimal-comma
    numbers ("1,3 ; 3,4", "0,8–1,9", "52,5–79,5") or two integers in a statistical context ("médiane 28 [24–33]").
    Parser v3: a decimal-comma pair is an interval when it stands alone (a table cell, the start of a line) or
    when its citation reading is not a well-formed list (see plausible_citation_list); the stat-context, value-
    before and two-decimal cues of v2 still apply."""
    tokens = re.findall(r"\d+", inner)
    if any(t.startswith("0") for t in tokens):          # [0,8–1,9] · [05] : never citation keys
        return True
    parts = [p for p in re.split(r"\s*[;–—-]\s*", inner.strip()) if p]
    if len(parts) != 2:
        return False
    stat = _stat_context(before)
    m = _NUMBER_JUST_BEFORE.search(before)
    after_value = bool(m) and not re.fullmatch(r"(19|20)\d{2}", m.group(1))
    standalone = not before.strip()                      # the bracket is the whole cell or opens the line
    dec = [re.fullmatch(r"\d+,\d+", p) for p in parts]
    if all(dec):
        sep = inner.strip()[len(parts[0]):].strip()[:1]
        if sep == ";":                                    # [1,3 ; 3,4] : French CI typography
            return True
        if stat or after_value or standalone or not plausible_citation_list(inner):
            return True
        return any(len(p.split(",")[1]) >= 2 for p in parts)
    if all(re.fullmatch(r"\d+", p) for p in parts):     # [24–33] : interval only with a statistical cue
        return stat or after_value or not plausible_citation_list(inner)
    return False


def _citation_brackets(text: str):
    """Yield the CITATION_NUMERIC matches that are citation keys (bracketed intervals are skipped)."""
    for m in CITATION_NUMERIC.finditer(text):
        if not bracket_is_interval(m.group(1), text[:m.start()]):
            yield m


def _blank(rx: re.Pattern, s: str) -> str:
    return rx.sub(lambda m: " " * len(m.group(0)), s)


def find_numbers(text: str, locale: str, strip_citations: bool = True) -> list[NumberHit]:
    """Return number hits in display order. Citation keys like [12] are blanked before extraction
    when strip_citations is True so that reference indices are not mistaken for statistics; bracketed
    intervals ([1,3 ; 3,4]) are kept. DOIs, URLs and multi-level section numbers (2.3.1) are ignored."""
    text = nfc(text).translate(ARABIC_INDIC)
    work = _blank(URL_DOI, text)
    work = _blank(HEX_HASH, work)
    work = _blank(MULTI_DOT, work)
    if strip_citations:
        spans = [(m.start(), m.end()) for m in _citation_brackets(work)]
        for a, b in reversed(spans):
            work = work[:a] + " " * (b - a) + work[b:]
    pat = _PATTERNS.get(locale, _PATTERNS["en"])
    hits: list[NumberHit] = []
    date_spans: list[tuple[int, int]] = []
    for m in DATE_NUMERIC.finditer(work):
        y, mo, d = (m.group(1), m.group(2), m.group(3)) if m.group(1) else (m.group(6), m.group(5), m.group(4))
        if not (1 <= int(d) <= 31 and 1 <= int(mo) <= 12 and 1900 <= int(y) <= 2100):
            continue              # "45-28-0001" is a catalogue number, not a date (v0.2.8): its parts are read as numbers
        norm = f"{int(d)}-{int(mo)}-{y}"
        ctx = text[max(0, m.start() - 60): m.end() + 60].replace("\n", " ")
        hits.append(NumberHit(raw=m.group(0), normalized=norm, start=m.start(), end=m.end(), role_hint="date", context=ctx))
        date_spans.append((m.start(), m.end()))
    for a, b in reversed(date_spans):
        work = work[:a] + " " * (b - a) + work[b:]
    for m in DATE_WORDS.finditer(work):
        gd = m.groupdict()
        mo = next(int(k[1:]) for k, v in gd.items() if v and k[0] in "mn" and k[1:].isdigit())
        d, y = (gd["d1"], gd["y1"]) if gd.get("d1") else (gd["d2"], gd["y2"])
        norm = f"{int(d)}-{mo}-{y}"
        ctx = text[max(0, m.start() - 60): m.end() + 60].replace("\n", " ")
        hits.append(NumberHit(raw=m.group(0), normalized=norm, start=m.start(), end=m.end(), role_hint="date", context=ctx))
    work = _blank(DATE_WORDS, work)
    pos = 0
    while True:
        m = pat.search(work, pos)
        if not m:
            break
        start, end, raw = m.start("num"), m.end("num"), m.group("num")
        before = work[max(0, start - 14): start]
        tail = m.end()  # end of the match including an ordinal suffix, if any
        if locale == "fr" and LABEL_BEFORE.search(before) and re.search(f"[{GROUP_SPACES}]", raw):
            # "tableau 3 120 femmes": the label number is 3, the value 120 is read on the next pass
            raw = re.match(r"[-−]?\d+", raw).group(0)
            end = tail = start + len(raw)
        pos = tail
        norm = normalize_number(raw, locale)
        after = work[tail: tail + 4]
        unit = None
        if re.match(rf"[{GROUP_SPACES}]?[%٪]", after):
            unit = "%"
        role = None
        if re.fullmatch(r"(19|20)\d{2}", norm) and unit is None:
            role = "year"
        if LABEL_BEFORE.search(before):
            role = "label_ref"
        elif re.search(r"(?i)\bp\s*[=<>≤≥]\s*$", before):
            role = "p_value"
        elif re.search(r"(?i)\bn\s*=\s*$", before):
            role = "n"
        elif not work[:start].strip() and re.match(r"[.)](?:\s|$)", work[end:end + 2]) and "." not in norm and "-" not in norm:
            # "1. …" / "2) …" opening a paragraph: list numbering, not a value of the paper (ignored on both sides)
            role = "numbering"
        ctx = text[max(0, start - 60): end + 60].replace("\n", " ")
        hits.append(NumberHit(raw=raw, normalized=norm, start=start, end=end, unit=unit, role_hint=role, context=ctx))
    hits.sort(key=lambda h: h.start)
    return hits


def number_multiset(text: str, locale: str, strip_citations: bool = True) -> dict[str, int]:
    out: dict[str, int] = {}
    for h in find_numbers(text, locale, strip_citations):
        out[h.normalized] = out.get(h.normalized, 0) + 1
    return out


# ---------------------------------------------------------------------------------------------
# Citations
# ---------------------------------------------------------------------------------------------
def expand_numeric_keys(inner: str) -> list[str]:
    keys: list[str] = []
    for part in re.split(r"\s*[,;]\s*", inner):
        part = part.strip()
        if not part:
            continue
        m = re.fullmatch(r"(\d{1,3})\s*[–—-]\s*(\d{1,3})", part)
        if m:
            a, b = int(m.group(1)), int(m.group(2))
            if a <= b and b - a < 60:
                keys.extend(str(i) for i in range(a, b + 1))
                continue
        if re.fullmatch(r"\d{1,3}", part):
            keys.append(part)
    return keys


def find_citations(text: str) -> dict[str, int]:
    text = nfc(text)
    counts: dict[str, int] = {}
    for m in _citation_brackets(text):
        for k in expand_numeric_keys(m.group(1)):
            key = f"[{k}]"
            counts[key] = counts.get(key, 0) + 1
    for m in CITATION_AUTHOR_DATE.finditer(text):
        for part in re.split(r"\s*;\s*", m.group(1)):
            key = citation_key(part)
            counts[key] = counts.get(key, 0) + 1
    for m in CITATION_NARRATIVE.finditer(text):
        key = citation_key(m.group(1) + ", " + m.group(2))
        counts[key] = counts.get(key, 0) + 1
    return counts


def citation_key(part: str) -> str:
    """One key for the same citation however it is written: "(Aiken and Clarke, 2002)", "Aiken & Clarke (2002)"
    and "Aiken et Clarke (2002)" all give "(Aiken & Clarke, 2002)"; "et al", "et al." and "وآخرون" give "et al."."""
    s = re.sub(r"\s+", " ", part.strip())
    s = _CIT_ETAL.sub(" et al.", s)
    s = _CIT_AND.sub(" & ", s)
    s = re.sub(r"\s*,?\s*((?:19|20)\d{2}[a-z]?)$", r", \1", s)
    return "(" + s.strip() + ")"


# ---------------------------------------------------------------------------------------------
# Document walking (python-docx)
# ---------------------------------------------------------------------------------------------
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _local(tag: str) -> str:
    return tag.rsplit('}', 1)[-1]


def _textbox_paragraphs(p_el):
    """Paragraph elements of the text boxes anchored in a paragraph (Word writes each box twice, as a DrawingML
    choice and a VML fallback: only the choice is read)."""
    for el in p_el.iter():
        if _local(el.tag) != "txbxContent":
            continue
        if any(_local(a.tag) == "Fallback" for a in el.iterancestors()):
            continue
        for child in el.iterchildren():
            if _local(child.tag) == "p":
                yield child


_NAV_GALLERIES = ("table of contents", "table des mati", "list of figures", "table of figures", "list of tables",
                  "table des illustrations", "table des figures", "liste des")
_NAV_STYLES = ("toc", "tm ", "tm1", "tm2", "tm3", "tm4", "table of figures", "table des illustrations",
               "en-tête de table", "en-ttedetable", "tableoffigures", "tableofcontents", "contents")


def _is_navigation_sdt(sdt_el) -> bool:
    """A table of contents / list of figures content control: navigation, not content (v0.2.7)."""
    for el in sdt_el.iter():
        if _local(el.tag) == "docPartGallery":
            val = (el.get("{%s}val" % W_NS) or "").lower()
            return any(k in val for k in _NAV_GALLERIES)
    return False


def is_navigation_paragraph(par) -> bool:
    """A paragraph in a TOC / list-of-figures style ("TM 2", "toc 1", "Table of Figures")."""
    st = par.style
    if st is None:
        return False
    name = (st.name or "").lower()
    sid = (getattr(st, "style_id", "") or "").lower()
    return name.startswith(_NAV_STYLES) or sid.startswith(_NAV_STYLES)


def iter_block_items(doc):
    """Yield ('p', Paragraph) and ('t', Table) in body order — including the content of body-level content
    controls (w:sdt, e.g. a bibliography inserted by Word's citation manager) and of text boxes; a table of
    contents or a list of figures (content control or TOC-styled paragraphs) is skipped (v0.2.7)."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    def walk(container):
        for child in container.iterchildren():
            tag = _local(child.tag)
            if tag == "p":
                par = Paragraph(child, doc)
                if is_navigation_paragraph(par):
                    continue
                yield "p", par
                for tb in _textbox_paragraphs(child):
                    yield "p", Paragraph(tb, doc)
            elif tag == "tbl":
                yield "t", Table(child, doc)
            elif tag == "sdt":
                if _is_navigation_sdt(child):
                    continue
                for sub in child.iterchildren():
                    if _local(sub.tag) == "sdtContent":
                        yield from walk(sub)

    yield from walk(doc.element.body)


_SKIP_ANCESTORS = {"del", "txbxContent", "Fallback", "instrText", "delText"}
SECTION_NUMBER = re.compile(r"^\s*\d{1,2}(?:\.\d{1,2}){1,3}\.?\s+[^\W\d_]")   # "2.3 Analyse statistique", "3.1.2. Mesures" — not "3.4 % des"


_SUP_KEY = re.compile(r"^\s*\d{1,3}(?:\s*[,;–—-]\s*\d{1,3})*\s*$")
_UNIT_BEFORE_SUP = re.compile(r"(?:^|[^A-Za-zÀ-ÿ])(?:k?m|cm|mm|dm|µm|m)\s?$")   # "kg/m" + superscript 2


def _is_superscript(t_el) -> bool:
    """True when the w:t sits in a run raised as a superscript."""
    for a in t_el.iterancestors():
        if _local(a.tag) == "r":
            for rpr in a:
                if _local(rpr.tag) == "rPr":
                    for prop in rpr:
                        if _local(prop.tag) == "vertAlign":
                            return prop.get(f"{{{W_NS}}}val") == "superscript"
            return False
    return False


def para_text(p_el, sup_cites: bool = False) -> str:
    """Text of a paragraph element as Word shows it: every w:t in document order, including hyperlinks, inline
    content controls (Word's citation fields) and tracked insertions; tracked deletions, field codes and the text
    boxes anchored in the paragraph (read separately) are left out. Tabs and line breaks become \\t and \\n.
    With sup_cites, a numeric superscript ("stages.⁴˒⁵" typed as a raised run "4,5") is read as the citation key
    "[4,5]" — the Vancouver superscript style — except after a unit ("kg/m" + raised 2) (v0.2.7)."""
    segs: list[tuple[str, bool]] = []
    for el in p_el.iter():
        tag = _local(el.tag)
        if tag not in ("t", "tab", "br", "cr", "noBreakHyphen", "softHyphen", "sym"):
            continue
        skip = False
        for a in el.iterancestors():
            if a is p_el:
                break
            if _local(a.tag) in _SKIP_ANCESTORS:
                skip = True
                break
        if skip:
            continue
        if tag == "t":
            segs.append((el.text or "", sup_cites and _is_superscript(el)))
        elif tag == "tab":
            segs.append(("\t", False))
        elif tag in ("br", "cr"):
            segs.append(("\n", False))
        elif tag == "noBreakHyphen":
            segs.append(("-", False))
    if not sup_cites:
        return "".join(s for s, _ in segs)
    out: list[str] = []
    i = 0
    while i < len(segs):
        if not segs[i][1]:
            out.append(segs[i][0])
            i += 1
            continue
        j = i
        while j < len(segs) and segs[j][1]:
            j += 1
        raised = "".join(s for s, _ in segs[i:j])
        before = "".join(out)
        if _SUP_KEY.match(raised) and not _UNIT_BEFORE_SUP.search(before) and not re.search(r"\d$", before):
            out.append("[" + re.sub(r"\s+", "", raised).replace(";", ",") + "]")
        else:
            out.append(raised)
        i = j
    return "".join(out)


def superscript_citations(doc) -> bool:
    """True when the document cites with raised numbers: at least three numeric superscript runs in its body."""
    n = 0
    for r in doc.element.body.iter():
        if _local(r.tag) != "r":
            continue
        rpr = next((c for c in r if _local(c.tag) == "rPr"), None)
        if rpr is None or not any(_local(x.tag) == "vertAlign" and x.get(f"{{{W_NS}}}val") == "superscript" for x in rpr):
            continue
        txt = "".join((t.text or "") for t in r if _local(t.tag) == "t")
        if re.fullmatch(r"\s*\d{1,3}(?:\s*[,;–—-]\s*\d{1,3})*\s*", txt):
            n += 1
            if n >= 3:
                return True
    return False


_SPLIT_THOUSANDS = re.compile(r"(?<![\d,.])(\d{1,3})\n(\d{3})(?!\d)")


def cell_text(tc_el, sup_cites: bool = False) -> str:
    """Text of a table cell: its paragraphs (nested tables included) joined by line breaks. A figure broken over a
    line break at a thousands boundary ("1⏎588", "500⏎000+") is one figure (v0.2.8): the break stands for the
    grouping space the student typed Enter instead of."""
    text = "\n".join(para_text(p, sup_cites) for p in tc_el.iter() if _local(p.tag) == "p")
    return _SPLIT_THOUSANDS.sub(r"\1 \2", text)


_TITLE_LABEL = re.compile(r"^\s*(?:chapitre|chapter|partie|part|section|[ée]tape|step|phase|axe|axis|objectif|objective|"
                          r"hypoth[èe]se|hypothesis|composante|component|الفصل|الباب|الجزء|المحور|الهدف|المرحلة|الخطوة)"
                          r"\s+(?:\d{1,2}|[IVX]{1,5})\s*[:.\-–—]?\s*", re.I)


def _carries_statistics(text: str) -> bool:
    """A line that reads like a result or a caption rather than a title: a figure that is not a section number or a
    title label, in a long line or next to a statistic sign ("Cinq dimensions présentent un α < 0,70 : …",
    "Tableau 2. Résultats des tests (N = 50)")."""
    rest = re.sub(r"^\s*\d{1,2}(?:\.\d{1,2}){0,3}[.)]?\s*", "", text.strip())
    rest = _TITLE_LABEL.sub("", rest)
    if not re.search(r"\d", rest):
        return False
    return len(text.strip()) > 90 or bool(re.search(r"[%=<>±≤≥]|\bp\s*[<=>]|\b[nN]\s*=|\(\s*\d|\d\s*\)", rest))


def is_heading(par) -> bool:
    name = (par.style.name if par.style is not None else "") or ""
    if name.lower().startswith(("heading", "titre", "title")):
        return True
    txt = para_text(par._p).strip()
    if not txt or len(txt) > 90 or "\n" in txt:
        return False
    numbered = bool(SECTION_NUMBER.match(txt))
    # a bold line that carries a figure ("Tableau 1 : Caractéristiques (N = 50)", "Le taux était de 85 %") is a
    # caption or a sentence, not a heading: its numbers belong to the ledger (v0.2.7)
    rest = re.sub(r"^\s*\d{1,2}(?:\.\d{1,2}){0,3}[.)]?\s*", "", txt)
    rest = _TITLE_LABEL.sub("", rest)            # "Chapitre 2 :", "Objectif 1 —", "Partie II" are still titles
    if re.search(r"\d", rest):
        return False
    runs = [r for r in par.runs if r.text.strip()]
    if runs and all(r.bold for r in runs) and (numbered or not txt.endswith((".", ";", ",", ":"))):
        return True
    # "2.3 Analyse statistique" typed as plain text: a multi-level section number opening a short line is a heading
    return numbered and not txt.endswith((".", ";", ",", ":"))


def docx_text_blocks(path: str) -> list[dict]:
    """Flatten a .docx into ordered blocks: {'kind': 'heading'|'paragraph'|'table', 'text', 'section', 'cells'?}."""
    import docx
    doc = docx.Document(path)
    blocks: list[dict] = []
    section = ""
    sup = superscript_citations(doc)
    for kind, item in iter_block_items(doc):
        if kind == "p":
            txt = nfc(para_text(item._p, sup))
            if not txt.strip():
                continue
            t = txt.strip()
            heading = is_heading(item)
            style_name = item.style.name if item.style else ""
            if heading and "\n" in t:
                # a heading style that carries body text after a soft line break ("2. PROBLEMATIQUE⏎Un dossier…"):
                # the first line is the heading, the rest is a paragraph whose numbers must reach the ledger (v0.2.7)
                head, _, rest = t.partition("\n")
                if head.strip() and rest.strip():
                    section = head.strip()
                    blocks.append({"kind": "heading", "text": section, "section": section, "style": style_name})
                    blocks.append({"kind": "paragraph", "text": rest.strip(), "section": section})
                    continue
            # The same line must be read the same way whatever its style, so that the ledger (French source) and
            # the target built from Markdown agree — the order below is deliberate (v0.2.7):
            if not heading and len(t) <= 60 and (is_refs_heading(t) or HEADING_ANNEX.match(t)) and not t.endswith((".", ";", ",")):
                heading = True        # "Références" / "RÉFÉRENCES" / "Annexes" typed as a plain line
            if heading and style_name.lower() not in ("title", "titre"):
                if len(t) > 160 and re.search(r"[.!?]\s+[A-ZÀ-ÝŒ]", t):
                    heading = False   # several sentences in a heading style: a paragraph
                elif _carries_statistics(t):
                    heading = False   # a result or a caption ("Annexe F : Structure du questionnaire (29 items)"):
                    #                   its figures must reach the ledger, whether typed bold, styled or plain
                elif len(t) >= 50 and is_reference_entry(t) and re.search(r"(19|20)\d{2}|10\.\d{4,9}/", t):
                    heading = False   # a reference entry typed in a heading style
            if heading:
                section = t
                blocks.append({"kind": "heading", "text": t, "section": section, "style": style_name})
            else:
                blocks.append({"kind": "paragraph", "text": txt, "section": section})
        else:
            cells = []
            seen_tc: list = []  # strong references: lxml returns the same proxy for a node while one is alive
            for row in item.rows:
                out_row = []
                for c in row.cells:
                    # a merged cell is returned once per grid column (and once per row for vertical merges):
                    # its text is read once, the other slots are empty (v0.2.7)
                    tc = c._tc
                    if any(tc is x for x in seen_tc):
                        out_row.append("")
                        continue
                    seen_tc.append(tc)
                    out_row.append(nfc(cell_text(tc, sup)).strip())
                cells.append(out_row)
            blocks.append({"kind": "table", "text": "\n".join("\t".join(r) for r in cells), "section": section, "cells": cells})
    return blocks


def is_md_rule_row(row: list[str]) -> bool:
    """True for a Markdown table rule line (|---|:--:|); an all-empty row is a real (empty) row."""
    cells = [c for c in row if c]
    return bool(cells) and all(re.fullmatch(r":?-{2,}:?", c) for c in cells)


def nonempty_rows(rows: list[list[str]]) -> int:
    """Number of table rows that carry at least one non-blank cell (spacer rows are not compared)."""
    return sum(1 for r in rows if any(str(c).strip() for c in r))


def docx_plain_text(path: str) -> str:
    return "\n".join(b["text"] for b in docx_text_blocks(path))


# ---------------------------------------------------------------------------------------------
# Markdown targets and the reference-list walk shared by every script (v0.2.7)
# ---------------------------------------------------------------------------------------------
_MD_LIST = re.compile(r"^\s*([-*•]|\d+[.)])\s+")


def md_blocks(src: str) -> list[dict]:
    """Flatten a Markdown (or the .txt dump of extract_ledger) into the same blocks as docx_text_blocks."""
    blocks: list[dict] = []
    lines = src.splitlines()
    i = 0
    section = ""
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            i += 1
            continue
        if ln.startswith("#"):
            section = ln.lstrip("#").strip()
            blocks.append({"kind": "heading", "text": nfc(section), "section": section})
            i += 1
            continue
        if ln.lstrip().startswith("|"):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                row = [nfc(c.strip()) for c in lines[i].strip().strip("|").split("|")]
                if not is_md_rule_row(row):  # empty rows are kept (v0.2.4)
                    rows.append(row)
                i += 1
            blocks.append({"kind": "table", "cells": rows, "section": section,
                           "text": "\n".join("\t".join(r) for r in rows)})
            continue
        if _MD_LIST.match(ln):  # list item (reference entries are usually numbered)
            blocks.append({"kind": "paragraph", "text": nfc(ln.strip()), "section": section})
            i += 1
            continue
        para = [ln]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith("#") and not lines[i].lstrip().startswith("|") \
                and not _MD_LIST.match(lines[i]):
            para.append(lines[i].rstrip())
            i += 1
        blocks.append({"kind": "paragraph", "text": nfc("\n".join(para)), "section": section})
    return blocks


def text_blocks(path: str) -> list[dict]:
    """Blocks of a .docx, .md or .txt file."""
    ext = path.lower().rsplit(".", 1)[-1]
    if ext == "docx":
        return docx_text_blocks(path)
    if ext in ("md", "txt", "markdown"):
        with open(path, encoding="utf-8") as f:
            return md_blocks(f.read())
    raise SystemExit(f"unsupported file: {path}")


def looks_like_reference(text: str) -> bool:
    """A reference-list entry: an index, or a year with the punctuation of a citation, or a DOI/URL, or 'et al.'."""
    t = strip_bullet(text)
    if not t or AFTER_REFS_SECTION.match(t):     # "Annexe A. Questionnaire" has the shape "Surname A." but is a title
        return False
    if REF_INDEX.match(t) or REF_AUTHOR_START.search(t[:60]):
        return True
    if re.search(r"(?<!\d)(19|20)\d{2}(?!\d)", t) and re.search(r"[;:,.]", t):
        return True
    if re.search(r"(?i)\bdoi\b|https?://|\b10\.\d{4,9}/", t):
        return True
    return bool(re.search(r"(?i)\bet\s+al\b", t))


def _split_on_line_breaks(entries: list[str]) -> list[str]:
    """A whole list typed in one paragraph with soft line breaks (Shift+Enter) is split back into entries: at
    every line that opens with an index when the paragraph is indexed, at every line that opens like an entry
    otherwise. A line that opens neither way is a wrapped line of the entry above (v0.2.7)."""
    out: list[str] = []
    for e in entries:
        if "\n" not in e:
            out.append(e)
            continue
        lines = [ln for ln in e.split("\n") if ln.strip()]
        if not lines:
            continue
        indexed = bool(REF_INDEX.match(strip_bullet(lines[0])))
        cur = lines[0]
        for ln in lines[1:]:
            s = strip_bullet(ln)
            new_entry = bool(REF_INDEX.match(s)) if indexed else is_reference_entry(s)
            if new_entry:
                out.append(cur)
                cur = ln
            else:
                cur = cur + " " + ln
        out.append(cur)
    return out


def merge_reference_fragments(entries: list[str]) -> list[str]:
    """Join a reference entry split over two paragraphs (Enter typed inside an entry): a fragment that carries no
    index while the previous entry had one, or that opens with a digit, a lower-case letter or punctuation, is
    appended to the previous entry (v0.2.7)."""
    merged: list[str] = []
    prev_indexed = False
    entries = [strip_bullet(e) for e in _split_on_line_breaks(entries)]
    entries = [e for e in entries if e]
    # an APA list carries no indices: there, "20 études, 12 pays." is an annotation of the entry above, not entry 20
    list_indexed = bool(entries) and bool(REF_INDEX.match(entries[0]))
    prev_idx: int | None = None
    for e in entries:
        im = REF_INDEX.match(e) if list_indexed else None
        indexed = im is not None
        # an entry starts with its authors ("Soufi G,", "von Elm E,", "Ministère de la Santé.") or an index; a
        # fragment ("2010;10:149.", "Rabat: MSPS; 2024.") shows neither in its first 60 characters
        authors = bool(REF_AUTHOR_START.search(e[:60]))
        if indexed and prev_idx is not None and not authors:
            idx = int(im.group(1))
            if idx not in (prev_idx, prev_idx + 1):
                indexed = False           # "20 études, 12 pays." after entry 7: an annotation, not entry 20
        opens_low = bool(re.match(r"[\d\W]|[a-zà-ÿ]", e))
        fragment = not indexed and not authors and (opens_low or not looks_like_reference(e))
        if list_indexed and not indexed and merged:
            fragment = True               # in an indexed list every entry carries its index: a line without one
            #                               is the wrapped rest of the entry above ("Evaluation of evidence. CMAJ…")
        if len(e) <= 120 and not indexed and not authors and re.search(r"\d{1,4}\s*\(\d{1,3}\)|\d+\s*[–-]\s*\d+\.?$", e):
            fragment = True                   # a journal/volume/pages line: "Reading Research Quarterly, 58(2), 285–312."
        if merged and ((prev_indexed and not indexed and not authors) or fragment):
            merged[-1] = merged[-1] + " " + e
            continue
        merged.append(e)
        prev_indexed = indexed
        if indexed:
            prev_idx = int(im.group(1))
    return merged


def is_reference_entry(text: str) -> bool:
    """A paragraph that OPENS like an entry: an index or an author/organisation (a year alone is not enough —
    prose after the list may well carry one)."""
    t = strip_bullet(text)
    return bool(t) and not AFTER_REFS_SECTION.match(t) and bool(REF_INDEX.match(t) or REF_AUTHOR_START.search(t[:60]))


def _refs_continue(blocks: list[dict], i: int, lookahead: int = 2) -> bool:
    """True when one of the next paragraphs opens like a reference entry (a wrapped entry is followed by more
    references; an annex title is followed by prose or a table)."""
    seen = 0
    for b in blocks[i + 1:]:
        if b["kind"] != "paragraph":
            return False
        if is_reference_entry(b["text"]):
            return True
        seen += 1
        if seen >= lookahead:
            return False
    return False


def walk_blocks(blocks: list[dict]):
    """Yield (block, in_refs): in_refs is True for the entries of the reference list only. The list starts at a
    "Références"-type heading and ends at the next heading, at a table, or at the first paragraph that is not a
    reference entry (an "Annexe A" title or an acknowledgement typed without a heading style, v0.2.7)."""
    in_refs = False
    entries_seen = 0
    for i, b in enumerate(blocks):
        if b["kind"] == "heading":
            in_refs = is_refs_heading(b["text"])
            entries_seen = 0
            yield b, False
            continue
        if b["kind"] == "table":
            in_refs = False
            yield b, False
            continue
        if in_refs:
            t = b["text"]
            if not looks_like_reference(t):
                if AFTER_REFS_SECTION.match(t) or not _refs_continue(blocks, i):
                    in_refs = False           # the list has ended (annex title, acknowledgement, prose)
                elif not entries_seen or NOTE_OPENER.match(t):
                    yield b, False            # a note ("Références conservées dans leur langue d'origine.",
                    continue                  # "Note : …") is body text, not an entry
                # otherwise: the tail of a wrapped entry ("Frontiers in Education.", "70(5), 285–312.") — kept
                # with the list and joined to its entry by merge_reference_fragments
            else:
                entries_seen += 1
        yield b, in_refs


def split_body_and_refs(blocks: list[dict]) -> tuple[str, list[list[list[str]]], list[str]]:
    """(body text, tables, reference entries) of a document: the body is every non-reference paragraph and every
    table (table/figure labels blanked so "Tableau 1" is not read as a statistic); headings are excluded."""
    body_parts, tables, refs = [], [], []
    for b, in_refs in walk_blocks(blocks):
        if b["kind"] == "heading":
            continue
        if b["kind"] == "table":
            tables.append(b["cells"])
            body_parts.append(b["text"])
        elif in_refs:
            refs.append(b["text"].strip())
        else:
            t = TABLE_LABEL.sub(lambda m: " " * len(m.group(0)), b["text"])
            t = FIGURE_LABEL.sub(lambda m: " " * len(m.group(0)), t)
            body_parts.append(t)
    return "\n".join(body_parts), tables, merge_reference_fragments(refs)


# ---------------------------------------------------------------------------------------------
# PPTX reading
# ---------------------------------------------------------------------------------------------
def pptx_slide_texts(path: str, include_notes: bool = False) -> list[dict]:
    from pptx import Presentation
    from pptx.util import Pt  # noqa: F401
    prs = Presentation(path)
    out = []
    for idx, slide in enumerate(prs.slides, 1):
        texts: list[str] = []
        tables: list[list[list[str]]] = []
        chart_values: list[dict] = []
        for shape in slide.shapes:
            if shape.is_placeholder and shape.placeholder_format.type is not None and \
                    str(shape.placeholder_format.type).startswith(("SLIDE_NUMBER", "DATE", "FOOTER")):
                continue
            if shape.has_text_frame:
                t = nfc(shape.text_frame.text)
                if t.strip():
                    texts.append(t)
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[nfc(c.text).strip() for c in r.cells] for r in shape.table.rows]
                tables.append(rows)
                texts.append("\n".join("\t".join(r) for r in rows))
            if getattr(shape, "has_chart", False) and shape.has_chart:
                try:
                    ch = shape.chart
                    cats = [str(c) for c in ch.plots[0].categories]
                    for s in ch.plots[0].series:
                        for c, v in zip(cats, s.values):
                            if v is None:
                                continue
                            chart_values.append({"series": s.name or "", "category": c,
                                                 "normalized": normalize_number(repr(float(v)).rstrip("0").rstrip(".") if isinstance(v, float) else str(v), "en")})
                        texts.append("CHART " + (s.name or "") + ": " + ", ".join(cats))
                except Exception:
                    pass
        notes = ""
        if include_notes and slide.has_notes_slide:
            notes = nfc(slide.notes_slide.notes_text_frame.text)
        out.append({"slide": idx, "layout": slide.slide_layout.name, "text": "\n".join(texts), "tables": tables,
                    "chart_values": chart_values, "notes": notes})
    return out


def save_json(obj, path: str) -> None:
    with open(path, "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False, indent=2)


def load_json(path: str):
    with open(path, encoding="utf-8") as f:
        return json.load(f)
